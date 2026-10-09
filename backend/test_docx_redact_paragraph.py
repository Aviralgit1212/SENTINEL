from __future__ import annotations

from docx import Document

from redact.docx_redact import (
    build_run_offset_table,
    redact_paragraph_span,
)


def test_single_run_redaction() -> None:
    document = Document()

    paragraph = document.add_paragraph(
        "Call me at 555-123-4567 today."
    )

    text, _ = build_run_offset_table(
        paragraph
    )

    start = text.index(
        "555-123-4567"
    )

    end = (
        start
        + len("555-123-4567")
    )

    result = redact_paragraph_span(
        paragraph,
        start,
        end,
    )

    assert result == (
        "Call me at [REDACTED] today."
    )

    assert (
        "555-123-4567"
        not in result
    )


def test_split_run_redaction() -> None:
    document = Document()

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "Call me at "
    )

    paragraph.add_run(
        "555-"
    )

    paragraph.add_run(
        "123-"
    )

    paragraph.add_run(
        "4567"
    )

    paragraph.add_run(
        " today."
    )

    text, spans = build_run_offset_table(
        paragraph
    )

    assert text == (
        "Call me at 555-123-4567 today."
    )

    assert len(spans) == 5

    start = text.index(
        "555-123-4567"
    )

    end = (
        start
        + len("555-123-4567")
    )

    result = redact_paragraph_span(
        paragraph,
        start,
        end,
    )

    assert result == (
        "Call me at [REDACTED] today."
    )

    assert (
        "555-123-4567"
        not in result
    )

    assert paragraph.runs[0].text == (
        "Call me at "
    )

    assert paragraph.runs[1].text == (
        "[REDACTED]"
    )

    assert paragraph.runs[2].text == ""

    assert paragraph.runs[3].text == ""

    assert paragraph.runs[4].text == (
        " today."
    )


def test_partial_span_inside_single_run() -> None:
    document = Document()

    paragraph = document.add_paragraph(
        "Account: 555-123-4567."
    )

    text, _ = build_run_offset_table(
        paragraph
    )

    number_start = text.index(
        "555-123-4567"
    )

    # Redact only "123" inside the single run.
    start = number_start + 4
    end = start + 3

    result = redact_paragraph_span(
        paragraph,
        start,
        end,
    )

    assert result == (
        "Account: 555-[REDACTED]-4567."
    )

    assert (
        "123"
        not in result[number_start:]
    )


def test_invalid_empty_span_rejected() -> None:
    document = Document()

    paragraph = document.add_paragraph(
        "Test"
    )

    try:
        redact_paragraph_span(
            paragraph,
            2,
            2,
        )
    except ValueError:
        pass
    else:
        raise AssertionError(
            "Empty entity span was not rejected."
        )


def main() -> None:
    test_single_run_redaction()

    print(
        "6.3.2 single-run redaction: PASS"
    )

    test_split_run_redaction()

    print(
        "6.3.2 split-run redaction: PASS"
    )

    test_partial_span_inside_single_run()

    print(
        "6.3.2 partial-run redaction: PASS"
    )

    test_invalid_empty_span_rejected()

    print(
        "6.3.2 invalid-span handling: PASS"
    )


if __name__ == "__main__":
    main()