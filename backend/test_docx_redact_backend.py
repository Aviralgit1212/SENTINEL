from __future__ import annotations

import io

from docx import Document

from extraction.extraction import extract_text
from redact.docx_redact import (
    redact_docx,
)


PHONE = "555-123-4567"


def build_test_document() -> bytes:
    document = Document()

    # Body.
    document.add_paragraph(
        f"Body contact: {PHONE}"
    )

    # Split-run body.
    split = document.add_paragraph()

    split.add_run(
        "Split contact: "
    )

    split.add_run(
        "555-"
    )

    split.add_run(
        "123-"
    )

    split.add_run(
        "4567"
    )

    # Table.
    table = document.add_table(
        rows=1,
        cols=1,
    )

    table.cell(
        0,
        0,
    ).text = (
        f"Table contact: {PHONE}"
    )

    # Header.
    section = document.sections[0]

    section.header.paragraphs[0].text = (
        f"Header contact: {PHONE}"
    )

    # Footer.
    section.footer.paragraphs[0].text = (
        f"Footer contact: {PHONE}"
    )

    output = io.BytesIO()

    document.save(
        output
    )

    return output.getvalue()


def main() -> None:
    original_bytes = build_test_document()

    original_text = extract_text(
        original_bytes,
        "AI_Guardian_DOCX_Redact_Test.docx",
    )

    assert original_text.count(
        PHONE
    ) == 5

    flagged_entities = []

    search_from = 0

    while True:
        start = original_text.find(
            PHONE,
            search_from,
        )

        if start == -1:
            break

        end = (
            start
            + len(PHONE)
        )

        flagged_entities.append(
            {
                "entity_type": "PHONE_NUMBER",
                "text": PHONE,
                "start": start,
                "end": end,
                "sensitive": True,
                "label": "Phone Number",
            }
        )

        search_from = end

    assert len(
        flagged_entities
    ) == 5

    redacted_bytes = redact_docx(
        original_bytes,
        flagged_entities,
    )

    assert redacted_bytes

    assert redacted_bytes != (
        original_bytes
    )

    # Final output must still be a valid DOCX.
    redacted_document = Document(
        io.BytesIO(
            redacted_bytes
        )
    )

    all_text: list[str] = []

    all_text.extend(
        paragraph.text
        for paragraph in redacted_document.paragraphs
    )

    for table in redacted_document.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text.extend(
                    paragraph.text
                    for paragraph in cell.paragraphs
                )

    for section in redacted_document.sections:
        all_text.extend(
            paragraph.text
            for paragraph in section.header.paragraphs
        )

        all_text.extend(
            paragraph.text
            for paragraph in section.footer.paragraphs
        )

    final_text = "\n".join(
        all_text
    )

    assert PHONE not in final_text

    assert final_text.count(
        "[REDACTED]"
    ) >= 5

    print(
        "6.3.5 DOCX backend REDACT integration: PASS"
    )

    print(
        "6.3.5 five sensitive locations redacted: PASS"
    )

    print(
        "6.3.5 generated DOCX validity: PASS"
    )


if __name__ == "__main__":
    main()