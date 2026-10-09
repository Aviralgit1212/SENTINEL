from __future__ import annotations

import io

from docx import Document

from redact.docx_redact import (
    AddressableParagraph,
    ParagraphRedactionTarget,
    iter_addressable_paragraphs,
    redact_document_targets,
)


PHONE = "555-123-4567"


def build_test_document() -> bytes:
    document = Document()

    # Body paragraph.
    document.add_paragraph(
        f"Body contact number: {PHONE}"
    )

    # Table cell.
    table = document.add_table(
        rows=1,
        cols=1,
    )

    table.cell(
        0,
        0,
    ).text = (
        f"Table contact number: {PHONE}"
    )

    # Header.
    section = document.sections[0]

    section.header.paragraphs[0].text = (
        f"Header contact number: {PHONE}"
    )

    # Footer.
    section.footer.paragraphs[0].text = (
        f"Footer contact number: {PHONE}"
    )

    # Second section.
    # Default header/footer remain linked to the previous section,
    # allowing the shared-part deduplication behavior to be tested.
    document.add_section()

    output = io.BytesIO()

    document.save(output)

    return output.getvalue()


def find_phone_span(
    paragraph_text: str,
) -> tuple[int, int]:
    start = paragraph_text.index(
        PHONE
    )

    end = (
        start
        + len(PHONE)
    )

    return (
        start,
        end,
    )


def test_all_required_sources_are_addressable() -> None:
    file_bytes = build_test_document()

    document = Document(
        io.BytesIO(file_bytes)
    )

    addressable = iter_addressable_paragraphs(
        document
    )

    sources = {
        item.source
        for item in addressable
    }

    assert "paragraph" in sources
    assert "table_cell" in sources
    assert "header" in sources
    assert "footer" in sources


def test_table_cell_has_paragraph_addressing() -> None:
    file_bytes = build_test_document()

    document = Document(
        io.BytesIO(file_bytes)
    )

    addressable = iter_addressable_paragraphs(
        document
    )

    table_entries = [
        item
        for item in addressable
        if item.source == "table_cell"
    ]

    assert len(table_entries) == 1

    entry = table_entries[0]

    assert entry.table_index == 0
    assert entry.row_index == 0
    assert entry.cell_index == 0
    assert entry.paragraph_index == 0

    assert PHONE in entry.paragraph.text


def test_shared_header_footer_parts_are_deduplicated() -> None:
    file_bytes = build_test_document()

    document = Document(
        io.BytesIO(file_bytes)
    )

    addressable = iter_addressable_paragraphs(
        document
    )

    headers = [
        item
        for item in addressable
        if item.source == "header"
    ]

    footers = [
        item
        for item in addressable
        if item.source == "footer"
    ]

    # The document has two sections, but the second section inherits the
    # first section's default header/footer package parts.
    assert len(headers) == 1
    assert len(footers) == 1

    assert headers[0].part_key is not None
    assert footers[0].part_key is not None


def test_redact_body_table_header_footer() -> None:
    file_bytes = build_test_document()

    document = Document(
        io.BytesIO(file_bytes)
    )

    addressable = iter_addressable_paragraphs(
        document
    )

    targets: list[
        ParagraphRedactionTarget
    ] = []

    for item in addressable:

        if PHONE not in item.paragraph.text:
            continue

        start, end = find_phone_span(
            item.paragraph.text
        )

        targets.append(
            ParagraphRedactionTarget(
                address=item,
                entity_start=start,
                entity_end=end,
            )
        )

    assert len(targets) == 4

    applied = redact_document_targets(
        document,
        targets,
    )

    assert applied == 4

    # Body.
    assert PHONE not in (
        document.paragraphs[0].text
    )

    # Table.
    assert PHONE not in (
        document.tables[0].cell(0, 0).text
    )

    # Shared header.
    assert PHONE not in (
        document.sections[0].header.paragraphs[0].text
    )

    # Shared footer.
    assert PHONE not in (
        document.sections[0].footer.paragraphs[0].text
    )


def test_redaction_preserves_non_sensitive_text() -> None:
    file_bytes = build_test_document()

    document = Document(
        io.BytesIO(file_bytes)
    )

    addressable = iter_addressable_paragraphs(
        document
    )

    target_item = next(
        item
        for item in addressable
        if item.source == "paragraph"
    )

    start, end = find_phone_span(
        target_item.paragraph.text
    )

    applied = redact_document_targets(
        document,
        [
            ParagraphRedactionTarget(
                address=target_item,
                entity_start=start,
                entity_end=end,
            )
        ],
    )

    assert applied == 1

    assert document.paragraphs[0].text == (
        "Body contact number: [REDACTED]"
    )


def main() -> None:
    test_all_required_sources_are_addressable()

    print(
        "6.3.3 body/table/header/footer addressing: PASS"
    )

    test_table_cell_has_paragraph_addressing()

    print(
        "6.3.3 table paragraph addressing: PASS"
    )

    test_shared_header_footer_parts_are_deduplicated()

    print(
        "6.3.3 shared header/footer deduplication: PASS"
    )

    test_redact_body_table_header_footer()

    print(
        "6.3.3 document-wide redaction application: PASS"
    )

    test_redaction_preserves_non_sensitive_text()

    print(
        "6.3.3 non-sensitive text preservation: PASS"
    )


if __name__ == "__main__":
    main()