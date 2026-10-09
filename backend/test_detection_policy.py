from __future__ import annotations

import base64
import io

import fitz
from docx import Document
from fastapi.testclient import TestClient

from main import app, run_security_pipeline, scanner

# Deterministically generated solely for this test (Verhoeff-valid; not a real record).
SYNTHETIC_AADHAAR = "234567890009"
SYNTHETIC_PAN = "ABCDE1234F"
SYNTHETIC_IFSC = "HDFC0001234"
SYNTHETIC_KEY = "sk-" + ("a7" * 14)


def _types(text: str) -> set[str]:
    return {entity["entity_type"] for entity in scanner.scan(text)}


def test_synthetic_indian_identifiers_and_credential_detection():
    text = f"Aadhaar candidate {SYNTHETIC_AADHAAR}; PAN {SYNTHETIC_PAN}; IFSC {SYNTHETIC_IFSC}; key {SYNTHETIC_KEY}"
    detected = _types(text)
    assert {"IN_AADHAAR", "IN_PAN", "IN_IFSC", "CREDENTIAL_TOKEN"}.issubset(detected)


def test_aadhaar_ocr_confusables_are_normalized_and_checksum_validated():
    confusable = SYNTHETIC_AADHAAR.replace("0", "O")
    assert "IN_AADHAAR" in _types(f"OCR candidate {confusable}")
    invalid = SYNTHETIC_AADHAAR[:-1] + ("8" if SYNTHETIC_AADHAAR[-1] != "8" else "7")
    assert "IN_AADHAAR" not in _types(f"Invalid candidate {invalid}")


def test_negative_lookalikes_do_not_match_custom_patterns():
    detected = _types("ABCDE12345F HDFC12345678 sk-short")
    assert not {"IN_AADHAAR", "IN_PAN", "IN_IFSC", "CREDENTIAL_TOKEN"}.intersection(detected)


def test_credential_risk_decision_is_critical_block():
    result = run_security_pipeline(f"test key {SYNTHETIC_KEY}")
    assert result["risk"] == "CRITICAL"
    assert result["decision"] == "BLOCK"


def test_sensitive_docx_core_property_is_selectively_redacted_and_verified():
    source_email = "synthetic.metadata.person@example.com"
    document = Document()
    document.add_paragraph("Synthetic document body remains unchanged.")
    document.core_properties.author = f"Example author {source_email}"
    document.core_properties.title = "Synthetic quarterly report"
    file_bytes = io.BytesIO()
    document.save(file_bytes)

    client = TestClient(app)
    result = client.post(
        "/scan-file",
        files={"file": ("metadata-test.docx", file_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert result.status_code == 200
    payload = result.json()
    assert payload["decision"] == "REDACT"
    assert payload["coverage"]["metadata_findings"] >= 1
    redacted_bytes = base64.b64decode(payload["redacted_file_base64"])
    redacted = Document(io.BytesIO(redacted_bytes))
    assert source_email not in redacted.core_properties.author
    assert "Synthetic quarterly report" == redacted.core_properties.title
    assert "Synthetic document body remains unchanged." in "\n".join(p.text for p in redacted.paragraphs)


def test_sensitive_pdf_info_metadata_is_selectively_redacted_and_verified():
    source_email = "synthetic.pdf.author@example.com"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Synthetic PDF body remains unchanged.")
    document.set_metadata({
        "title": "Synthetic preserved title", "author": source_email,
        "subject": "Synthetic test subject", "keywords": "",
        "creator": "Synthetic creator", "producer": "Synthetic producer",
        "creationDate": "", "modDate": "", "trapped": "",
    })
    source = document.tobytes()
    document.close()

    response = TestClient(app).post(
        "/scan-file",
        files={"file": ("metadata-test.pdf", source, "application/pdf")},
        headers={"X-Sentinel-Cache-Choice": "rescan"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["decision"] == "REDACT"
    assert payload["coverage"]["metadata_findings"] >= 1
    assert payload["redaction_status"] == "VERIFIED_SUPPORTED_TEXT_AND_INFO_METADATA"
    redacted = fitz.open(stream=base64.b64decode(payload["redacted_file_base64"]), filetype="pdf")
    try:
        assert source_email not in (redacted.metadata.get("author") or "")
        assert redacted.metadata["title"] == "Synthetic preserved title"
        assert "Synthetic PDF body remains unchanged." in "\n".join(page.get_text() for page in redacted)
    finally:
        redacted.close()
