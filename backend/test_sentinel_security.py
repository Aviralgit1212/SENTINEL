from __future__ import annotations

import asyncio
import base64
import io
import json
import sqlite3
from pathlib import Path

import fitz
import pytest
from fastapi import HTTPException
from starlette.datastructures import UploadFile
from docx import Document
from PIL import Image

import main
import cache_store
import extraction.extraction as extraction
from extraction.extraction import inspect_docx_coverage
from dashboard_store import DB_PATH, create_security_event, update_event_state
import dashboard_store
from release_store import (
    IDLE_TIMEOUT_SECONDS,
    clear_held_payloads,
    cleanup_expired_payloads,
    consume_release,
    create_held_payload,
    memory_status,
)


class _Request:
    headers = {"X-Sentinel-Site": "https://example.test"}


def _create_block_event() -> str:
    return create_security_event(
        username="test-reviewer", site="example.test", input_type="TEXT",
        filename=None, risk="HIGH", decision="BLOCK", state="BLOCKED",
        entities=[{"entity_type": "US_SSN", "text": "synthetic-test-value"}],
    )


def test_audit_database_stores_only_normalized_filename():
    event_id, _ = main._record_dashboard_event(
        request=_Request(), input_type="PDF", filename="private-ssn-123456789.pdf",
        security_result={"decision": "ALLOW", "risk": "LOW", "entities": []},
    )
    event = dashboard_store.get_security_event(event_id)
    assert event["filename"] == "uploaded.pdf"
    assert "123456789" not in json.dumps(event)


def test_original_payload_is_process_memory_only_and_manual_clear_works():
    clear_held_payloads()
    event_id = _create_block_event()
    token = "one-time-test-token"
    create_held_payload(event_id=event_id, release_token=token, payload_type="text", raw_text="synthetic-sensitive-value")
    assert memory_status()["items"] == 1
    assert memory_status()["bytes"] == len("synthetic-sensitive-value")
    with sqlite3.connect(DB_PATH) as db:
        tables = {row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert "held_payloads" not in tables
    assert clear_held_payloads() == 1
    assert memory_status() == {"items": 0, "bytes": 0}


def test_idle_cleanup_drops_expired_memory_references(monkeypatch):
    clear_held_payloads()
    event_id = _create_block_event()
    create_held_payload(event_id=event_id, release_token="expiry-token", payload_type="text", raw_text="transient")
    monkeypatch.setattr("release_store.IDLE_TIMEOUT_SECONDS", 0)
    assert cleanup_expired_payloads() == 1
    assert memory_status() == {"items": 0, "bytes": 0}


def test_approved_release_is_one_time_and_original_not_persisted():
    clear_held_payloads()
    event_id = _create_block_event()
    token = "single-use-token"
    create_held_payload(event_id=event_id, release_token=token, payload_type="text", raw_text="transient-original")
    update_event_state(event_id=event_id, new_state="APPROVED", actor="test-reviewer", reviewed_by="test-reviewer")
    released = consume_release(event_id=event_id, release_token=token)
    assert released["raw_text"] == "transient-original"
    assert memory_status() == {"items": 0, "bytes": 0}
    with pytest.raises(PermissionError):
        consume_release(event_id=event_id, release_token=token)


def test_upload_size_limit_uses_actual_bytes(monkeypatch):
    monkeypatch.setattr(main, "MAX_FILE_BYTES", 4)
    upload = UploadFile(filename="small.pdf", file=io.BytesIO(b"12345"))
    with pytest.raises(HTTPException) as error:
        asyncio.run(main.scan_file(_Request(), upload))
    assert error.value.status_code == 413


def test_xlsx_is_rejected_without_entering_scanner():
    upload = UploadFile(filename="sheet.xlsx", file=io.BytesIO(b"PK"))
    with pytest.raises(HTTPException) as error:
        asyncio.run(main.scan_file(_Request(), upload))
    assert error.value.status_code == 415


def test_docx_embedded_images_are_reported_as_uninspected():
    image_buffer = io.BytesIO()
    Image.new("RGB", (4, 4), "white").save(image_buffer, format="PNG")
    doc = Document()
    doc.add_paragraph("Synthetic body text")
    doc.add_picture(io.BytesIO(image_buffer.getvalue()))
    output = io.BytesIO()
    doc.save(output)
    coverage = inspect_docx_coverage(output.getvalue())
    assert "embedded images/media" in coverage["unsupported_layers"]


def test_docx_sensitive_core_property_is_scannable():
    doc = Document()
    doc.add_paragraph("Synthetic body text")
    doc.core_properties.author = "synthetic@example.com"
    output = io.BytesIO()
    doc.save(output)
    coverage = inspect_docx_coverage(output.getvalue())
    assert "author" in coverage["metadata_fields_present"]
    result = main.run_security_pipeline(coverage["metadata_text"])
    assert result["entities"]


def test_empty_ocr_result_is_incomplete_not_allow(monkeypatch):
    class _OCR:
        full_text = ""
        regions = []
    monkeypatch.setattr(main, "extract_image", lambda _: _OCR())
    monkeypatch.setattr(main, "run_security_pipeline", lambda _: {
        "entities": [], "risk": "LOW", "decision": "ALLOW", "redacted_text": None,
    })
    upload = UploadFile(filename="blank.png", file=io.BytesIO(b"image-bytes"), headers={"content-type": "image/png"})
    result = asyncio.run(main.scan_file(_Request(), upload))
    assert result["decision"] == "INCOMPLETE"
    assert result["scan_status"] == "INCOMPLETE"
    assert result["redacted_file_base64"] is None
    assert cache_store.get_cached_report(cache_store.sha256_fingerprint(b"image-bytes")) is None


def test_pdf_ocr_processing_stops_at_fifteen_pages(monkeypatch):
    from PIL import Image
    png = io.BytesIO()
    Image.new("RGB", (2, 2), "white").save(png, format="PNG")
    png_bytes = png.getvalue()
    document = fitz.open()
    for _ in range(16):
        page = document.new_page(width=100, height=100)
        page.insert_image(page.rect, stream=png_bytes)
    source = document.tobytes()
    document.close()
    class _Region:
        confidence = 90.0
    class _OCR:
        full_text = "safe synthetic text"
        regions = [_Region()]
    calls = {"count": 0}
    def fake_ocr(_):
        calls["count"] += 1
        return _OCR()
    monkeypatch.setattr(extraction, "extract_image", fake_ocr)
    result = extraction.extract_pdf_analysis(source)
    assert calls["count"] == 15
    assert result["scanned_page_count"] == 16
    assert result["complete"] is False
    assert any("limit" in reason for reason in result["incomplete_reasons"])


def test_pdf_ocr_without_confidence_is_incomplete(monkeypatch):
    png = io.BytesIO()
    Image.new("RGB", (4, 4), "white").save(png, format="PNG")
    document = fitz.open()
    page = document.new_page(width=100, height=100)
    page.insert_image(page.rect, stream=png.getvalue())
    source = document.tobytes()
    document.close()

    class _Region:
        confidence = -1.0

    class _OCR:
        full_text = "safe synthetic OCR text"
        regions = [_Region()]

    monkeypatch.setattr(extraction, "extract_image", lambda _image: _OCR())
    result = extraction.extract_pdf_analysis(source)
    assert result["complete"] is False
    assert result["pages"][0]["ocr_status"] == "uncertain"
    assert any("uncertain" in reason for reason in result["incomplete_reasons"])


def test_extension_manifest_contains_gemini_and_no_gmail():
    root = Path(__file__).resolve().parents[1]
    manifest = json.loads((root / "extension" / "manifest.json").read_text())
    text = (root / "extension" / "content.js").read_text(encoding="utf-8")
    assert any("https://*/*" == match for script in manifest["content_scripts"] for match in script["matches"])
    assert "gemini.google.com" in text
    assert "gmail.com" not in text.lower()
    assert "Sentinel protection may not be active here" in text
    assert "sessionStorage" in text
    assert "Sentinel could not complete this scan" in text
    assert "not released to the website" in text
