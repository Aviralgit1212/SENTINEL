from __future__ import annotations

import io

from docx import Document

from extraction.extraction import (
    extract_docx_elements,
    extract_text,
)


def build_test_docx() -> bytes:
    document = Document()

    section = document.sections[0]

    section.header.paragraphs[0].text = (
        "HEADER_MARKER_TEXT"
    )

    section.footer.paragraphs[0].text = (
        "FOOTER_MARKER_TEXT"
    )

    document.add_paragraph(
        "PARAGRAPH_MARKER_TEXT"
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(
        0,
        0,
    ).text = "CELL_R1C1"

    table.cell(
        0,
        1,
    ).text = "CELL_R1C2"

    table.cell(
        1,
        0,
    ).text = "CELL_R2C1"

    table.cell(
        1,
        1,
    ).text = "CELL_R2C2"

    output = io.BytesIO()

    document.save(output)

    return output.getvalue()


def main() -> None:
    file_bytes = build_test_docx()

    text = extract_text(
        file_bytes,
        "AI_Guardian_Test.docx",
    )

    required_markers = [
        "PARAGRAPH_MARKER_TEXT",
        "CELL_R1C1",
        "CELL_R1C2",
        "CELL_R2C1",
        "CELL_R2C2",
        "HEADER_MARKER_TEXT",
        "FOOTER_MARKER_TEXT",
    ]

    for marker in required_markers:
        assert marker in text, (
            f"Missing extracted marker: {marker}"
        )

    elements = extract_docx_elements(
        file_bytes
    )

    element_types = {
        element["type"]
        for element in elements
    }

    assert "paragraph" in element_types
    assert "table_cell" in element_types
    assert "header" in element_types
    assert "footer" in element_types

    # Corrupted DOCX must fail.
    try:
        extract_text(
            b"this is not a docx file",
            "broken.docx",
        )
    except ValueError:
        pass
    else:
        raise AssertionError(
            "Corrupted DOCX did not raise ValueError."
        )

    # Unsupported extension must fail.
    try:
        extract_text(
            b"anything",
            "test.txt",
        )
    except ValueError:
        pass
    else:
        raise AssertionError(
            "Unsupported extension did not raise ValueError."
        )

    print(
        "DOCX extraction test: PASS"
    )

    print(
        "Addressable DOCX elements: PASS"
    )

    print(
        "Corrupted DOCX test: PASS"
    )

    print(
        "Unsupported extension test: PASS"
    )


if __name__ == "__main__":
    main()