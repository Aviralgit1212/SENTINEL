from __future__ import annotations

import io
import re
import zipfile
import xml.etree.ElementTree as ET
import fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from extraction.image_ocr import extract_image


PDF_OCR_PAGE_LIMIT = 15


def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extract text from supported document formats.

    Supported:
        - PDF
        - DOCX

    Returns:
        All extractable text as one string.

    Raises:
        ValueError: If the file type is unsupported or the document
            cannot be opened/read.
    """
    filename_lower = (filename or "").lower()

    # --------------------------------------------------
    # PDF — EXISTING BEHAVIOR PRESERVED
    # --------------------------------------------------
    if filename_lower.endswith(".pdf"):
        doc = None

        try:
            # Open directly from memory; no temporary file is written to disk.
            doc = fitz.open(
                stream=file_bytes,
                filetype="pdf",
            )

            # A password-protected PDF may open successfully but still require
            # a password before its contents can be read.
            if doc.needs_pass:
                raise ValueError(
                    "Could not extract text from a password-protected PDF."
                )

            page_text = []

            for page in doc:
                page_text.append(page.get_text())

            extracted_text = "\n".join(page_text)

            # KNOWN LIMITATION:
            # If this PDF is image-only/scanned and has no text layer,
            # page.get_text() may return empty text.
            # OCR/Tesseract support is planned for a later step.
            return extracted_text

        except ValueError:
            raise

        except Exception as exc:
            raise ValueError(
                "Could not extract text from PDF. "
                "The file may be corrupted, unreadable, or unsupported."
            ) from exc

        finally:
            if doc is not None:
                doc.close()

    # --------------------------------------------------
    # DOCX
    # --------------------------------------------------
    if filename_lower.endswith(".docx"):
        try:
            document = Document(
                io.BytesIO(file_bytes)
            )

            fragments: list[str] = []

            # Headers from every section.
            for section in document.sections:
                for paragraph in section.header.paragraphs:
                    text = paragraph.text

                    if text and text.strip():
                        fragments.append(text)

            # Main document paragraphs.
            for paragraph in document.paragraphs:
                text = paragraph.text

                if text and text.strip():
                    fragments.append(text)

            # Tables and their cells.
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text

                        if text and text.strip():
                            fragments.append(text)

            # Footers from every section.
            for section in document.sections:
                for paragraph in section.footer.paragraphs:
                    text = paragraph.text

                    if text and text.strip():
                        fragments.append(text)

            # Newlines deliberately separate independent document fragments.
            return "\n".join(fragments)

        except (
            PackageNotFoundError,
            KeyError,
            ValueError,
            OSError,
        ) as exc:
            raise ValueError(
                "Corrupted or unreadable DOCX file."
            ) from exc

        except Exception as exc:
            raise ValueError(
                "Corrupted or unreadable DOCX file."
            ) from exc

    # --------------------------------------------------
    # UNSUPPORTED FORMAT
    # --------------------------------------------------
    raise ValueError(
        f"Unsupported file type: {filename or 'unknown'}"
    )


def extract_pages(file_bytes: bytes) -> list[str]:
    """
    Extract selectable PDF text page-by-page, fully in memory.

    This function exists for PDF redaction only.
    """
    doc = None

    try:
        # Open directly from memory; no temporary file is written to disk.
        doc = fitz.open(
            stream=file_bytes,
            filetype="pdf",
        )

        if doc.needs_pass:
            raise ValueError(
                "Could not extract pages from a password-protected PDF."
            )

        return [
            page.get_text()
            for page in doc
        ]

    except ValueError:
        raise

    except Exception as exc:
        raise ValueError(
            "Could not extract PDF pages. "
            "The file may be corrupted, unreadable, or unsupported."
        ) from exc

    finally:
        if doc is not None:
            doc.close()


def extract_pdf_analysis(file_bytes: bytes) -> dict:
    """Inspect selectable PDF text and image-bearing pages in memory.

    Pages containing embedded images are rendered at approximately 300 DPI
    and OCR'd. Only the first 15 such pages are processed. Any omitted page,
    OCR failure, or unreadable image layer is reported as incomplete instead
    of being treated as clean.
    """
    doc = None
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        if doc.needs_pass:
            raise ValueError("Could not inspect a password-protected PDF.")
        fragments: list[str] = []
        pages: list[dict] = []
        ocr_ranges: list[tuple[int, int]] = []
        reasons: list[str] = []
        scanned_page_count = 0
        text_length = 0

        for page_number, page in enumerate(doc, start=1):
            page_text = page.get_text() or ""
            has_images = bool(page.get_images(full=True))
            page_record = {
                "page": page_number,
                "selectable_text": bool(page_text.strip()),
                "image_layer": has_images,
                "ocr_status": "not_required",
                "ocr_confidence": None,
            }
            page_parts = [page_text] if page_text else []

            if has_images:
                scanned_page_count += 1
                if scanned_page_count > PDF_OCR_PAGE_LIMIT:
                    page_record["ocr_status"] = "not_processed_limit"
                    reasons.append(f"Page {page_number} image layer was not processed: OCR limit is {PDF_OCR_PAGE_LIMIT} pages.")
                else:
                    try:
                        scale = 300 / 72
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                        ocr_result = extract_image(pixmap.tobytes("png"))
                        confidences = [r.confidence for r in ocr_result.regions if r.confidence >= 0]
                        average = (sum(confidences) / len(confidences)) if confidences else None
                        page_record["ocr_status"] = "complete" if ocr_result.full_text.strip() and average is not None and average >= 40 else "uncertain"
                        page_record["ocr_confidence"] = round(average, 2) if average is not None else None
                        if not ocr_result.full_text.strip() or average is None or average < 40:
                            reasons.append(f"Page {page_number} OCR coverage is uncertain.")
                        page_parts.append(ocr_result.full_text)
                    except Exception:
                        # Keep diagnostics generic: exceptions may contain source text.
                        page_record["ocr_status"] = "failed"
                        reasons.append(f"Page {page_number} OCR or rendering failed.")

            page_content = "\n".join(part for part in page_parts if part)
            if page_content:
                if fragments:
                    text_length += 1
                start = text_length
                fragments.append(page_content)
                text_length += len(page_content)
                if has_images and page_record["ocr_status"] in {"complete", "uncertain_empty"}:
                    # OCR text is the suffix after selectable text and its separator.
                    ocr_text = page_parts[-1] if page_parts else ""
                    if ocr_text:
                        ocr_start = start + (len(page_text) + (1 if page_text else 0))
                        ocr_ranges.append((ocr_start, ocr_start + len(ocr_text)))
            pages.append(page_record)

        metadata = doc.metadata or {}
        metadata_values = {
            field: str(metadata.get(field) or "")
            for field in ("title", "author", "subject", "keywords", "creator", "producer")
        }
        try:
            xml_metadata_present = bool((doc.get_xml_metadata() or "").strip())
        except Exception:
            xml_metadata_present = True
            reasons.append("PDF XML metadata could not be inspected reliably.")
        if xml_metadata_present:
            reasons.append("Uninspected PDF XML metadata is present.")

        return {
            "text": "\n".join(fragments),
            "pages": pages,
            "page_count": len(doc),
            "scanned_page_count": scanned_page_count,
            "ocr_page_limit": PDF_OCR_PAGE_LIMIT,
            "ocr_ranges": ocr_ranges,
            "incomplete_reasons": list(dict.fromkeys(reasons)),
            "complete": not reasons,
            "metadata_values": metadata_values,
            "metadata_fields_present": sorted(key for key, value in metadata_values.items() if value.strip()),
            "xml_metadata_present": xml_metadata_present,
        }
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not inspect PDF content safely.") from exc
    finally:
        if doc is not None:
            doc.close()


def extract_docx_elements(
    file_bytes: bytes,
) -> list[dict]:
    """
    Extract addressable DOCX text elements for future DOCX redaction.

    This function is read-only.

    Each returned element identifies where the text came from:
        - header
        - paragraph
        - table_cell
        - footer

    Table-cell entries are paragraph-granular so future redaction can
    deterministically identify the exact paragraph within a cell.
    """
    try:
        document = Document(
            io.BytesIO(file_bytes)
        )

        elements: list[dict] = []

        # --------------------------------------------------
        # HEADERS
        # --------------------------------------------------
        for section_index, section in enumerate(
            document.sections
        ):
            for paragraph_index, paragraph in enumerate(
                section.header.paragraphs
            ):
                text = paragraph.text

                if text and text.strip():
                    elements.append(
                        {
                            "type": "header",
                            "section_index": section_index,
                            "paragraph_index": paragraph_index,
                            "text": text,
                        }
                    )

        # --------------------------------------------------
        # BODY PARAGRAPHS
        # --------------------------------------------------
        for paragraph_index, paragraph in enumerate(
            document.paragraphs
        ):
            text = paragraph.text

            if text and text.strip():
                elements.append(
                    {
                        "type": "paragraph",
                        "paragraph_index": paragraph_index,
                        "text": text,
                    }
                )

        # --------------------------------------------------
        # TABLE CELLS — PARAGRAPH GRANULAR
        # --------------------------------------------------
        for table_index, table in enumerate(
            document.tables
        ):
            for row_index, row in enumerate(
                table.rows
            ):
                for cell_index, cell in enumerate(
                    row.cells
                ):
                    for paragraph_index, paragraph in enumerate(
                        cell.paragraphs
                    ):
                        text = paragraph.text

                        if text and text.strip():
                            elements.append(
                                {
                                    "type": "table_cell",
                                    "table_index": table_index,
                                    "row_index": row_index,
                                    "cell_index": cell_index,
                                    "paragraph_index": paragraph_index,
                                    "text": text,
                                }
                            )

        # --------------------------------------------------
        # FOOTERS
        # --------------------------------------------------
        for section_index, section in enumerate(
            document.sections
        ):
            for paragraph_index, paragraph in enumerate(
                section.footer.paragraphs
            ):
                text = paragraph.text

                if text and text.strip():
                    elements.append(
                        {
                            "type": "footer",
                            "section_index": section_index,
                            "paragraph_index": paragraph_index,
                            "text": text,
                        }
                    )

        return elements

    except (
        PackageNotFoundError,
        KeyError,
        ValueError,
        OSError,
    ) as exc:
        raise ValueError(
            "Corrupted or unreadable DOCX file."
        ) from exc

    except Exception as exc:
        raise ValueError(
            "Corrupted or unreadable DOCX file."
        ) from exc


def inspect_docx_coverage(file_bytes: bytes) -> dict:
    """Identify DOCX layers the current extractor/redactor cannot verify.

    Only labels and transient core-property text are returned; callers must
    not persist the values. Unknown secondary structures are surfaced as
    incomplete rather than implicitly treated as inspected.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded DOCX package exceeds the inspection limit.")
            names = set(archive.namelist())
            unsupported: set[str] = set()
            doc_xml = archive.read("word/document.xml") if "word/document.xml" in names else b""
            if any(name.startswith("word/media/") for name in names):
                unsupported.add("embedded images/media")
            if "word/comments.xml" in names:
                unsupported.add("comments")
            if "word/footnotes.xml" in names or "word/endnotes.xml" in names:
                unsupported.add("footnotes/endnotes")
            if b"<w:txbxContent" in doc_xml:
                unsupported.add("text boxes")
            if b"<w:ins" in doc_xml or b"<w:del" in doc_xml:
                unsupported.add("tracked changes")
            if any(name.startswith("word/embeddings/") for name in names):
                unsupported.add("custom XML or embedded objects")
            for name in names:
                if not name.startswith("customXml/") or not name.endswith(".xml") or "/_rels/" in name or "itemProps" in name:
                    continue
                try:
                    custom_root = ET.fromstring(archive.read(name))
                    custom_text = "".join(custom_root.itertext()).strip()
                    has_data_children = any(True for _ in custom_root)
                    if custom_text or has_data_children:
                        unsupported.add("custom XML or embedded objects")
                except ET.ParseError:
                    unsupported.add("custom XML or embedded objects")
            table_count = len(re.findall(rb"<w:tbl(?:\s[^>]*)?>", doc_xml))
            document = Document(io.BytesIO(file_bytes))
            if table_count > len(document.tables):
                unsupported.add("nested tables")
            for name in names:
                if name.startswith(("word/header", "word/footer")) and name.endswith(".xml"):
                    part = archive.read(name)
                    if b"<w:tbl" in part:
                        unsupported.add("header/footer tables")
                if name.endswith(".rels") and b'TargetMode="External"' in archive.read(name):
                    unsupported.add("external relationship targets")

        properties = document.core_properties
        metadata = {
            "title": properties.title or "",
            "subject": properties.subject or "",
            "author": properties.author or "",
            "keywords": properties.keywords or "",
            "comments": properties.comments or "",
            "category": properties.category or "",
            "last_modified_by": properties.last_modified_by or "",
        }
        metadata_text = "\n".join(value for value in metadata.values() if value.strip())
        return {
            "unsupported_layers": sorted(unsupported),
            "metadata_fields_present": sorted(key for key, value in metadata.items() if value.strip()),
            "metadata_text": metadata_text,
            "metadata_values": metadata,
        }
    except Exception as exc:
        raise ValueError("Could not inspect DOCX package layers safely.") from exc
