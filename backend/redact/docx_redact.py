from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Iterable

from docx import Document
from docx.enum.section import WD_HEADER_FOOTER_INDEX
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from extraction.extraction import (
    extract_docx_elements,
    extract_text,
)


DEFAULT_REDACTION_PLACEHOLDER = "[REDACTED]"


class DocxRedactionError(RuntimeError):
    """Raised when DOCX redaction or verification fails."""


@dataclass(frozen=True)
class RunSpan:
    """
    Character span for one DOCX run inside a concatenated paragraph string.

    start is inclusive.
    end is exclusive.
    """

    run_index: int
    start: int
    end: int
    run: Run


@dataclass(frozen=True)
class AddressableParagraph:
    """
    A paragraph together with enough structural information to identify it
    deterministically inside a DOCX document.

    source:
        paragraph
        table_cell
        header
        footer
    """

    source: str
    paragraph: Paragraph
    paragraph_index: int

    section_index: int | None = None
    table_index: int | None = None
    row_index: int | None = None
    cell_index: int | None = None

    variant: str = "default"

    part_key: str | None = None


@dataclass(frozen=True)
class ParagraphRedactionTarget:
    """
    A security-approved span inside one addressable DOCX paragraph.

    Offsets refer to the concatenated run text of that paragraph.
    """

    address: AddressableParagraph
    entity_start: int
    entity_end: int


@dataclass(frozen=True)
class DocumentParagraphSpan:
    """
    A paragraph together with its position inside the same flattened
    document text contract used by extract_text().

    start is inclusive.
    end is exclusive.
    """

    address: AddressableParagraph
    start: int
    end: int


def build_run_offset_table(
    paragraph: Paragraph,
) -> tuple[str, list[RunSpan]]:
    """
    Build the exact concatenated paragraph text and a character-offset
    table mapping each DOCX run to its location inside that text.

    No document mutation occurs.
    """

    fragments: list[str] = []
    spans: list[RunSpan] = []

    cursor = 0

    for run_index, run in enumerate(
        paragraph.runs
    ):
        run_text = run.text or ""

        start = cursor
        end = start + len(run_text)

        fragments.append(run_text)

        spans.append(
            RunSpan(
                run_index=run_index,
                start=start,
                end=end,
                run=run,
            )
        )

        cursor = end

    return (
        "".join(fragments),
        spans,
    )


def find_overlapping_run_spans(
    run_spans: list[RunSpan],
    entity_start: int,
    entity_end: int,
) -> list[RunSpan]:
    """
    Return every run whose character span overlaps the supplied entity span.

    Both spans use half-open interval notation:
        [start, end)

    No document mutation occurs.
    """

    if entity_start < 0:
        raise ValueError(
            "Entity start offset cannot be negative."
        )

    if entity_end < entity_start:
        raise ValueError(
            "Entity end offset cannot be before start."
        )

    if entity_start == entity_end:
        return []

    overlapping: list[RunSpan] = []

    for run_span in run_spans:
        if (
            run_span.start < entity_end
            and entity_start < run_span.end
        ):
            overlapping.append(
                run_span
            )

    return overlapping


def redact_paragraph_span(
    paragraph: Paragraph,
    entity_start: int,
    entity_end: int,
    placeholder: str = DEFAULT_REDACTION_PLACEHOLDER,
) -> str:
    """
    Redact exactly one entity span inside one DOCX paragraph.

    The placeholder is inserted only into the first affected run.
    Additional affected runs have their sensitive text removed.

    Returns:
        Paragraph text after redaction.

    Raises:
        ValueError:
            If the supplied span is invalid or cannot be mapped.
    """

    paragraph_text, run_spans = build_run_offset_table(
        paragraph
    )

    paragraph_length = len(
        paragraph_text
    )

    if entity_start < 0:
        raise ValueError(
            "Entity start offset cannot be negative."
        )

    if entity_end < entity_start:
        raise ValueError(
            "Entity end offset cannot be before start."
        )

    if entity_start == entity_end:
        raise ValueError(
            "Entity span cannot be empty."
        )

    if entity_end > paragraph_length:
        raise ValueError(
            "Entity span exceeds paragraph text length."
        )

    if not isinstance(
        placeholder,
        str,
    ):
        raise ValueError(
            "Redaction placeholder must be a string."
        )

    overlapping = find_overlapping_run_spans(
        run_spans,
        entity_start,
        entity_end,
    )

    if not overlapping:
        raise ValueError(
            "Entity span does not overlap any DOCX run."
        )

    first_run_index = overlapping[0].run_index

    for run_span in overlapping:
        run_text = run_span.run.text or ""

        overlap_start = max(
            entity_start,
            run_span.start,
        )

        overlap_end = min(
            entity_end,
            run_span.end,
        )

        local_start = (
            overlap_start
            - run_span.start
        )

        local_end = (
            overlap_end
            - run_span.start
        )

        prefix = run_text[:local_start]
        suffix = run_text[local_end:]

        if (
            run_span.run_index
            == first_run_index
        ):
            run_span.run.text = (
                prefix
                + placeholder
                + suffix
            )
        else:
            run_span.run.text = (
                prefix
                + suffix
            )

    return "".join(
        run.text or ""
        for run in paragraph.runs
    )


def _get_part_key(
    container: object,
) -> str:
    """
    Return a stable identity for an underlying DOCX package part.

    Headers and footers can be linked across sections and therefore refer
    to the same underlying package part.
    """

    part = getattr(
        container,
        "part",
        None,
    )

    if part is not None:
        part_name = getattr(
            part,
            "partname",
            None,
        )

        if part_name is not None:
            return str(
                part_name
            )

        return f"object-part:{id(part)}"

    element = getattr(
        container,
        "_element",
        None,
    )

    if element is not None:
        return f"element-part:{id(element)}"

    return f"container:{id(container)}"


def _header_footer_variants(
    section: object,
) -> Iterable[tuple[str, object]]:
    """
    Yield only header/footer variants that are actually defined in the
    section XML.

    This avoids materializing phantom linked containers.
    """

    header_properties = {
        WD_HEADER_FOOTER_INDEX.PRIMARY: "header",
        WD_HEADER_FOOTER_INDEX.FIRST_PAGE: "first_page_header",
        WD_HEADER_FOOTER_INDEX.EVEN_PAGE: "even_page_header",
    }

    footer_properties = {
        WD_HEADER_FOOTER_INDEX.PRIMARY: "footer",
        WD_HEADER_FOOTER_INDEX.FIRST_PAGE: "first_page_footer",
        WD_HEADER_FOOTER_INDEX.EVEN_PAGE: "even_page_footer",
    }

    sect_pr = getattr(
        section,
        "_sectPr",
        None,
    )

    if sect_pr is None:
        return

    for reference in sect_pr.headerReference_lst:
        property_name = header_properties.get(
            reference.type_
        )

        if property_name is not None:
            yield (
                property_name,
                getattr(
                    section,
                    property_name,
                ),
            )

    for reference in sect_pr.footerReference_lst:
        property_name = footer_properties.get(
            reference.type_
        )

        if property_name is not None:
            yield (
                property_name,
                getattr(
                    section,
                    property_name,
                ),
            )


def iter_addressable_paragraphs(
    document: Document,
) -> list[AddressableParagraph]:
    """
    Enumerate every text-bearing DOCX paragraph that 6.3 can redact.

    Covers:
        - body paragraphs
        - table-cell paragraphs
        - explicitly defined headers
        - explicitly defined footers

    Shared header/footer package parts are yielded only once for mutation.
    """

    result: list[AddressableParagraph] = []

    # --------------------------------------------------
    # BODY PARAGRAPHS
    # --------------------------------------------------

    for paragraph_index, paragraph in enumerate(
        document.paragraphs
    ):
        if not paragraph.text:
            continue

        result.append(
            AddressableParagraph(
                source="paragraph",
                paragraph=paragraph,
                paragraph_index=paragraph_index,
            )
        )

    # --------------------------------------------------
    # TABLE CELL PARAGRAPHS
    # --------------------------------------------------

    for table_index, table in enumerate(
        document.tables
    ):
        for row_index, row in enumerate(
            table.rows
        ):
            for cell_index, cell in enumerate(
                row.cells
            ):
                for paragraph_index, paragraph in enumerate(
                    cell.paragraphs
                ):
                    if not paragraph.text:
                        continue

                    result.append(
                        AddressableParagraph(
                            source="table_cell",
                            paragraph=paragraph,
                            paragraph_index=paragraph_index,
                            table_index=table_index,
                            row_index=row_index,
                            cell_index=cell_index,
                        )
                    )

    # --------------------------------------------------
    # EXPLICITLY DEFINED HEADERS / FOOTERS
    # --------------------------------------------------

    processed_parts: set[str] = set()

    for section_index, section in enumerate(
        document.sections
    ):
        for variant_key, container in (
            _header_footer_variants(section)
        ):
            part_key = _get_part_key(
                container
            )

            if part_key in processed_parts:
                continue

            processed_parts.add(
                part_key
            )

            is_header = (
                "header"
                in variant_key
            )

            source = (
                "header"
                if is_header
                else "footer"
            )

            if "first_page" in variant_key:
                variant = "first_page"
            elif "even_page" in variant_key:
                variant = "even_page"
            else:
                variant = "default"

            paragraphs = getattr(
                container,
                "paragraphs",
                [],
            )

            for paragraph_index, paragraph in enumerate(
                paragraphs
            ):
                if not paragraph.text:
                    continue

                result.append(
                    AddressableParagraph(
                        source=source,
                        paragraph=paragraph,
                        paragraph_index=paragraph_index,
                        section_index=section_index,
                        variant=variant,
                        part_key=part_key,
                    )
                )

    return result


def _iter_extraction_paragraphs(
    document: Document,
) -> list[AddressableParagraph]:
    """
    Enumerate paragraphs in EXACTLY the same logical order as the current
    extract_docx_elements()/extract_text() DOCX contract.

    This intentionally does NOT deduplicate shared header/footer parts.

    Why:
        extract_text() can contain the same inherited header/footer text
        more than once across sections. Entity offsets are based on that
        flattened extraction string, so mapping must use the same sequence.

    Mutation itself later deduplicates the underlying physical paragraph.
    """

    result: list[AddressableParagraph] = []

    # --------------------------------------------------
    # HEADERS
    # --------------------------------------------------

    for section_index, section in enumerate(
        document.sections
    ):
        for paragraph_index, paragraph in enumerate(
            section.header.paragraphs
        ):
            if not paragraph.text:
                continue

            result.append(
                AddressableParagraph(
                    source="header",
                    paragraph=paragraph,
                    paragraph_index=paragraph_index,
                    section_index=section_index,
                    variant="default",
                )
            )

    # --------------------------------------------------
    # BODY
    # --------------------------------------------------

    for paragraph_index, paragraph in enumerate(
        document.paragraphs
    ):
        if not paragraph.text:
            continue

        result.append(
            AddressableParagraph(
                source="paragraph",
                paragraph=paragraph,
                paragraph_index=paragraph_index,
            )
        )

    # --------------------------------------------------
    # TABLES
    # --------------------------------------------------

    for table_index, table in enumerate(
        document.tables
    ):
        for row_index, row in enumerate(
            table.rows
        ):
            for cell_index, cell in enumerate(
                row.cells
            ):
                for paragraph_index, paragraph in enumerate(
                    cell.paragraphs
                ):
                    if not paragraph.text:
                        continue

                    result.append(
                        AddressableParagraph(
                            source="table_cell",
                            paragraph=paragraph,
                            paragraph_index=paragraph_index,
                            table_index=table_index,
                            row_index=row_index,
                            cell_index=cell_index,
                        )
                    )

    # --------------------------------------------------
    # FOOTERS
    # --------------------------------------------------

    for section_index, section in enumerate(
        document.sections
    ):
        for paragraph_index, paragraph in enumerate(
            section.footer.paragraphs
        ):
            if not paragraph.text:
                continue

            result.append(
                AddressableParagraph(
                    source="footer",
                    paragraph=paragraph,
                    paragraph_index=paragraph_index,
                    section_index=section_index,
                    variant="default",
                )
            )

    return result


def _build_document_paragraph_spans(
    document: Document,
) -> list[DocumentParagraphSpan]:
    """
    Build global document offsets using the same fragment ordering and
    newline joining semantics used by extract_text().

    Each paragraph span corresponds to one logical extracted fragment.
    """

    elements = _iter_extraction_paragraphs(
        document
    )

    spans: list[
        DocumentParagraphSpan
    ] = []

    cursor = 0

    for index, address in enumerate(
        elements
    ):
        text = address.paragraph.text

        start = cursor
        end = start + len(text)

        spans.append(
            DocumentParagraphSpan(
                address=address,
                start=start,
                end=end,
            )
        )

        cursor = end

        if index < len(elements) - 1:
            cursor += 1

    return spans


def _entity_target_from_global_span(
    paragraph_spans: list[DocumentParagraphSpan],
    entity_text: str,
    entity_start: int,
    entity_end: int,
) -> ParagraphRedactionTarget:
    """
    Resolve a scanner entity's global document offsets to exactly one DOCX
    paragraph and convert them to local paragraph offsets.

    No substring guessing across paragraphs is allowed.
    """

    if not isinstance(
        entity_text,
        str,
    ) or not entity_text:
        raise DocxRedactionError(
            "Invalid sensitive entity text."
        )

    if not isinstance(
        entity_start,
        int,
    ) or not isinstance(
        entity_end,
        int,
    ):
        raise DocxRedactionError(
            "Invalid sensitive entity offsets."
        )

    if entity_start < 0:
        raise DocxRedactionError(
            "Sensitive entity start offset is negative."
        )

    if entity_end <= entity_start:
        raise DocxRedactionError(
            "Sensitive entity span is empty or inverted."
        )

    containing = [
        span
        for span in paragraph_spans
        if (
            span.start
            <= entity_start
            and entity_end
            <= span.end
        )
    ]

    if len(containing) != 1:
        raise DocxRedactionError(
            "Sensitive entity does not map to exactly one DOCX paragraph."
        )

    paragraph_span = containing[0]

    local_start = (
        entity_start
        - paragraph_span.start
    )

    local_end = (
        entity_end
        - paragraph_span.start
    )

    paragraph_text = (
        paragraph_span.address.paragraph.text
    )

    extracted_entity = paragraph_text[
        local_start:local_end
    ]

    if extracted_entity != entity_text:
        raise DocxRedactionError(
            "Sensitive entity text does not match the mapped DOCX paragraph."
        )

    return ParagraphRedactionTarget(
        address=paragraph_span.address,
        entity_start=local_start,
        entity_end=local_end,
    )


def redact_document_targets(
    document: Document,
    targets: Iterable[ParagraphRedactionTarget],
    placeholder: str = DEFAULT_REDACTION_PLACEHOLDER,
) -> int:
    """
    Apply security-approved redaction spans across the complete DOCX.

    Targets belonging to the same physical paragraph are processed from
    right to left so original offsets remain valid.

    Shared header/footer paragraph wrappers are deduplicated by their
    underlying XML paragraph element.

    Returns:
        Number of unique physical redaction spans applied.
    """

    target_groups: dict[
        int,
        tuple[
            Paragraph,
            list[ParagraphRedactionTarget],
        ],
    ] = {}

    seen_target_keys: set[
        tuple[int, int, int]
    ] = set()

    for target in targets:
        if not isinstance(
            target,
            ParagraphRedactionTarget,
        ):
            raise DocxRedactionError(
                "Invalid DOCX redaction target."
            )

        paragraph = target.address.paragraph

        paragraph_element = getattr(
            paragraph,
            "_p",
            None,
        )

        if paragraph_element is None:
            paragraph_element = getattr(
                paragraph,
                "_element",
                None,
            )

        if paragraph_element is None:
            paragraph_element = paragraph

        physical_paragraph_id = id(
            paragraph_element
        )

        target_key = (
            physical_paragraph_id,
            target.entity_start,
            target.entity_end,
        )

        if target_key in seen_target_keys:
            continue

        seen_target_keys.add(
            target_key
        )

        if physical_paragraph_id not in target_groups:
            target_groups[
                physical_paragraph_id
            ] = (
                paragraph,
                [],
            )

        target_groups[
            physical_paragraph_id
        ][1].append(
            target
        )

    applied_count = 0

    for paragraph, paragraph_targets in (
        target_groups.values()
    ):
        paragraph_targets.sort(
            key=lambda target: (
                target.entity_start,
                target.entity_end,
            ),
            reverse=True,
        )

        for target in paragraph_targets:
            try:
                redact_paragraph_span(
                    paragraph,
                    target.entity_start,
                    target.entity_end,
                    placeholder=placeholder,
                )
            except Exception as exc:
                raise DocxRedactionError(
                    "Could not apply a DOCX paragraph redaction."
                ) from exc

            applied_count += 1

    return applied_count


def serialize_docx(
    document: Document,
) -> bytes:
    """
    Serialize a DOCX document fully in memory.
    """

    output = io.BytesIO()

    try:
        document.save(
            output
        )
    except Exception as exc:
        raise DocxRedactionError(
            "Could not generate redacted DOCX."
        ) from exc

    redacted_bytes = output.getvalue()

    if not redacted_bytes:
        raise DocxRedactionError(
            "Generated DOCX is empty."
        )

    return redacted_bytes


def verify_redacted_docx_bytes(
    redacted_bytes: bytes,
    original_sensitive_values: Iterable[str],
) -> str:
    """
    Re-open and verify generated DOCX bytes.

    Verification:
        1. Open with python-docx.
        2. Re-extract through the same extract_text() function.
        3. Confirm every original sensitive string is absent.
    """

    if not isinstance(
        redacted_bytes,
        (bytes, bytearray, memoryview),
    ):
        raise DocxRedactionError(
            "Generated DOCX bytes must be bytes-like data."
        )

    redacted_bytes = bytes(
        redacted_bytes
    )

    if not redacted_bytes:
        raise DocxRedactionError(
            "Generated DOCX is empty."
        )

    # --------------------------------------------------
    # REOPEN
    # --------------------------------------------------

    try:
        verification_document = Document(
            io.BytesIO(redacted_bytes)
        )
    except Exception as exc:
        raise DocxRedactionError(
            "Generated DOCX could not be reopened."
        ) from exc

    # --------------------------------------------------
    # RE-EXTRACT
    # --------------------------------------------------

    try:
        verification_text = extract_text(
            redacted_bytes,
            "redacted_verification.docx",
        )
        properties = verification_document.core_properties
        metadata_text = "\n".join(
            str(getattr(properties, field, "") or "")
            for field in (
                "title", "subject", "author", "keywords", "comments",
                "category", "last_modified_by",
            )
        )
        verification_text += "\n" + metadata_text
    except Exception as exc:
        raise DocxRedactionError(
            "Could not re-extract generated DOCX for verification."
        ) from exc

    # --------------------------------------------------
    # NORMALIZE ORIGINAL VALUES
    # --------------------------------------------------

    unique_sensitive_values: list[str] = []
    seen: set[str] = set()

    for value in original_sensitive_values:
        if not isinstance(
            value,
            str,
        ):
            continue

        if not value:
            continue

        if value in seen:
            continue

        seen.add(
            value
        )

        unique_sensitive_values.append(
            value
        )

    # --------------------------------------------------
    # VERIFY EVERY VALUE
    # --------------------------------------------------

    for sensitive_value in unique_sensitive_values:
        if sensitive_value in verification_text:
            raise DocxRedactionError(
                "Post-redaction verification found an original "
                "sensitive value in the generated DOCX."
            )

    return verification_text


def redact_and_verify_document(
    document: Document,
    original_sensitive_values: Iterable[str],
    placeholder: str = DEFAULT_REDACTION_PLACEHOLDER,
) -> bytes:
    """
    Serialize a modified DOCX and verify it before release.
    """

    redacted_bytes = serialize_docx(
        document
    )

    verify_redacted_docx_bytes(
        redacted_bytes,
        original_sensitive_values,
    )

    return redacted_bytes


def redact_docx(
    file_bytes: bytes,
    flagged_entities: Iterable[dict],
    placeholder: str = DEFAULT_REDACTION_PLACEHOLDER,
    metadata_entities: Iterable[dict] = (),
) -> bytes:
    """
    Perform complete DOCX REDACT for already security-approved entities.

    The Scanner / Policy / Risk / Decision pipeline is NOT run here.

    flagged_entities must contain, for each sensitive entity:
        {
            "text": str,
            "start": int,
            "end": int,
            ...
        }

    Process:
        DOCX bytes
        → reopen
        → build global paragraph map
        → map each flagged entity to exact paragraph/run offsets
        → redact
        → serialize
        → reopen
        → re-extract
        → verify every original sensitive value is absent

    Raises:
        DocxRedactionError:
            On any mapping, mutation, serialization, reopening, or
            verification failure.
    """

    if not isinstance(
        file_bytes,
        (bytes, bytearray, memoryview),
    ):
        raise DocxRedactionError(
            "DOCX input must be bytes-like data."
        )

    file_bytes = bytes(
        file_bytes
    )

    if not file_bytes:
        raise DocxRedactionError(
            "DOCX input is empty."
        )

    entities = list(flagged_entities)
    metadata_entities = list(metadata_entities)

    if not entities and not metadata_entities:
        raise DocxRedactionError(
            "No redaction targets were supplied for a REDACT decision."
        )

    # --------------------------------------------------
    # OPEN ORIGINAL DOCX
    # --------------------------------------------------

    try:
        document = Document(
            io.BytesIO(
                file_bytes
            )
        )
    except Exception as exc:
        raise DocxRedactionError(
            "Could not open DOCX for redaction."
        ) from exc

    # --------------------------------------------------
    # BUILD GLOBAL OFFSET MAP
    # --------------------------------------------------

    try:
        paragraph_spans = (
            _build_document_paragraph_spans(
                document
            )
        )

        # Also obtain the extraction layer's addressability representation.
        # The call ensures the redaction adapter is aligned with the same
        # extractor contract used elsewhere in the backend.
        extraction_elements = extract_docx_elements(
            file_bytes
        )

        extraction_text = extract_text(
            file_bytes,
            "original_verification.docx",
        )

        flattened_from_paragraphs = "\n".join(
            span.address.paragraph.text
            for span in paragraph_spans
        )

        if (
            flattened_from_paragraphs
            != extraction_text
        ):
            raise DocxRedactionError(
                "DOCX redaction mapping does not match the extraction contract."
            )

        if len(extraction_elements) != len(
            paragraph_spans
        ):
            raise DocxRedactionError(
                "DOCX addressability map is inconsistent with extracted elements."
            )

    except DocxRedactionError:
        raise

    except Exception as exc:
        raise DocxRedactionError(
            "Could not build DOCX redaction mapping."
        ) from exc

    # --------------------------------------------------
    # MAP ALL SECURITY-APPROVED ENTITIES
    # --------------------------------------------------

    targets: list[
        ParagraphRedactionTarget
    ] = []

    original_sensitive_values: list[str] = []

    for entity in entities:
        if not isinstance(
            entity,
            dict,
        ):
            raise DocxRedactionError(
                "Invalid security entity returned to DOCX redactor."
            )

        entity_text = entity.get(
            "text"
        )

        entity_start = entity.get(
            "start"
        )

        entity_end = entity.get(
            "end"
        )

        target = _entity_target_from_global_span(
            paragraph_spans,
            entity_text,
            entity_start,
            entity_end,
        )

        targets.append(
            target
        )

        original_sensitive_values.append(
            entity_text
        )

    # --------------------------------------------------
    # APPLY
    # --------------------------------------------------

    applied = redact_document_targets(
        document,
        targets,
        placeholder=placeholder,
    )

    if applied <= 0 and not metadata_entities:
        raise DocxRedactionError(
            "No DOCX redactions were applied."
        )

    # Core properties are separately addressable, so redact only the fields
    # and spans that the scanner actually identified. Never wipe all metadata.
    property_names = {
        "title", "subject", "author", "keywords", "comments",
        "category", "last_modified_by",
    }
    grouped_metadata: dict[str, list[dict]] = {}
    for entity in metadata_entities:
        field = entity.get("metadata_field")
        if field not in property_names:
            raise DocxRedactionError("Unsupported DOCX metadata field redaction target.")
        grouped_metadata.setdefault(field, []).append(entity)

    properties = document.core_properties
    for field, field_entities in grouped_metadata.items():
        value = str(getattr(properties, field, "") or "")
        spans: list[tuple[int, int, str]] = []
        for entity in field_entities:
            entity_text = entity.get("text")
            start, end = entity.get("start"), entity.get("end")
            if (
                not isinstance(entity_text, str) or not entity_text
                or not isinstance(start, int) or not isinstance(end, int)
                or start < 0 or end <= start or end > len(value)
                or value[start:end] != entity_text
            ):
                raise DocxRedactionError("DOCX metadata span did not match the inspected property value.")
            spans.append((start, end, entity_text))
            original_sensitive_values.append(entity_text)
        for start, end, _text in sorted(spans, reverse=True):
            value = value[:start] + placeholder + value[end:]
        setattr(properties, field, value)

    # --------------------------------------------------
    # SERIALIZE + VERIFY
    # --------------------------------------------------

    return redact_and_verify_document(
        document,
        original_sensitive_values,
        placeholder=placeholder,
    )
