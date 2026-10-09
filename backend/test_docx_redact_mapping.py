from __future__ import annotations

import io

from docx import Document

from extraction.extraction import (
    extract_docx_elements,
)
from redact.docx_redact import (
    build_run_offset_table,
    find_overlapping_run_spans,
)


def build_test_document() -> bytes:
    document = Document()

    document.add_paragraph(
        "ORDINARY_PARAGRAPH"
    )

    split_paragraph = document.add_paragraph()

    split_paragraph.add_run(
        "Call me at "
    )

    split_paragraph.add_run(
        "555-"
    )

    split_paragraph.add_run(
        "123-"
    )

    split_paragraph.add_run(
        "4567"
    )

    table = document.add_table(
        rows=1,
        cols=1,
    )

    cell = table.cell(
        0,
        0,
    )

    cell.paragraphs[0].text = (
        "TABLE_FIRST_PARAGRAPH"
    )

    cell.add_paragraph(
        "TABLE_SECOND_PARAGRAPH"
    )

    section = document.sections[0]

    section.header.paragraphs[0].text = (
        "HEADER_MARKER"
    )

    section.footer.paragraphs[0].text = (
        "FOOTER_MARKER"
    )

    output = io.BytesIO()

    document.save(output)

    return output.getvalue()


def test_run_offset_mapping() -> None:
    document = Document()

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "Call me at "
    )

    paragraph.add_run(
        "555-"
    )

    paragraph.add_run(
        "123-"
    )

    paragraph.add_run(
        "4567"
    )

    text, spans = build_run_offset_table(
        paragraph
    )

    assert text == (
        "Call me at 555-123-4567"
    )

    assert len(spans) == 4

    assert spans[0].start == 0
    assert spans[0].end == 11

    assert spans[1].start == 11
    assert spans[1].end == 15

    assert spans[2].start == 15
    assert spans[2].end == 19

    assert spans[3].start == 19
    assert spans[3].end == 23

    entity_start = text.index(
        "555-123-4567"
    )

    entity_end = (
        entity_start
        + len("555-123-4567")
    )

    overlapping = find_overlapping_run_spans(
        spans,
        entity_start,
        entity_end,
    )

    assert [
        span.run_index
        for span in overlapping
    ] == [1, 2, 3]


def test_table_cell_is_paragraph_addressable() -> None:
    file_bytes = build_test_document()

    elements = extract_docx_elements(
        file_bytes
    )

    table_elements = [
        element
        for element in elements
        if element["type"] == "table_cell"
    ]

    assert len(table_elements) == 2

    assert table_elements[0][
        "table_index"
    ] == 0

    assert table_elements[0][
        "row_index"
    ] == 0

    assert table_elements[0][
        "cell_index"
    ] == 0

    assert table_elements[0][
        "paragraph_index"
    ] == 0

    assert table_elements[0][
        "text"
    ] == "TABLE_FIRST_PARAGRAPH"

    assert table_elements[1][
        "paragraph_index"
    ] == 1

    assert table_elements[1][
        "text"
    ] == "TABLE_SECOND_PARAGRAPH"


def test_all_docx_sources_are_addressable() -> None:
    file_bytes = build_test_document()

    elements = extract_docx_elements(
        file_bytes
    )

    element_types = {
        element["type"]
        for element in elements
    }

    assert "paragraph" in element_types
    assert "table_cell" in element_types
    assert "header" in element_types
    assert "footer" in element_types


def main() -> None:
    test_run_offset_mapping()

    print(
        "6.3.1 run-offset mapping: PASS"
    )

    test_table_cell_is_paragraph_addressable()

    print(
        "6.3.1 table paragraph addressing: PASS"
    )

    test_all_docx_sources_are_addressable()

    print(
        "6.3.1 header/footer/body/table addressing: PASS"
    )


if __name__ == "__main__":
    main()