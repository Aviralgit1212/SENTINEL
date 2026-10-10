from __future__ import annotations

import io

import fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extract text from supported document formats.

    Supported:
        - PDF
        - DOCX

    Returns:
        All extractable text as one string.

    Raises:
        ValueError: If the file type is unsupported or the document
            cannot be opened/read.
    """
    filename_lower = (filename or "").lower()

    # --------------------------------------------------
    # PDF — EXISTING BEHAVIOR PRESERVED
    # --------------------------------------------------
    if filename_lower.endswith(".pdf"):
        doc = None

        try:
            # Open directly from memory; no temporary file is written to disk.
            doc = fitz.open(
                stream=file_bytes,
                filetype="pdf",
            )

            # A password-protected PDF may open successfully but still require
            # a password before its contents can be read.
            if doc.needs_pass:
                raise ValueError(
                    "Could not extract text from a password-protected PDF."
                )

            page_text = []

            for page in doc:
                page_text.append(page.get_text())

            extracted_text = "\n".join(page_text)

            # KNOWN LIMITATION:
            # If this PDF is image-only/scanned and has no text layer,
            # page.get_text() may return empty text.
            # OCR/Tesseract support is planned for a later step.
            return extracted_text

        except ValueError:
            raise

        except Exception as exc:
            raise ValueError(
                "Could not extract text from PDF. "
                "The file may be corrupted, unreadable, or unsupported."
            ) from exc

        finally:
            if doc is not None:
                doc.close()

    # --------------------------------------------------
    # DOCX
    # --------------------------------------------------
    if filename_lower.endswith(".docx"):
        try:
            document = Document(
                io.BytesIO(file_bytes)
            )

            fragments: list[str] = []

            # Headers from every section.
            for section in document.sections:
                for paragraph in section.header.paragraphs:
                    text = paragraph.text

                    if text and text.strip():
                        fragments.append(text)

            # Main document paragraphs.
            for paragraph in document.paragraphs:
                text = paragraph.text

                if text and text.strip():
                    fragments.append(text)

            # Tables and their cells.
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text

                        if text and text.strip():
                            fragments.append(text)

            # Footers from every section.
            for section in document.sections:
                for paragraph in section.footer.paragraphs:
                    text = paragraph.text

                    if text and text.strip():
                        fragments.append(text)

            # Newlines deliberately separate independent document fragments.
            return "\n".join(fragments)

        except (
            PackageNotFoundError,
            KeyError,
            ValueError,
            OSError,
        ) as exc:
            raise ValueError(
                "Corrupted or unreadable DOCX file."
            ) from exc

        except Exception as exc:
            raise ValueError(
                "Corrupted or unreadable DOCX file."
            ) from exc

    # --------------------------------------------------
    # UNSUPPORTED FORMAT
    # --------------------------------------------------
    raise ValueError(
        f"Unsupported file type: {filename or 'unknown'}"
    )


def extract_pages(file_bytes: bytes) -> list[str]:
    """
    Extract selectable PDF text page-by-page, fully in memory.

    This function exists for PDF redaction only.
    """
    doc = None

    try:
        # Open directly from memory; no temporary file is written to disk.
        doc = fitz.open(
            stream=file_bytes,
            filetype="pdf",
        )

        if doc.needs_pass:
            raise ValueError(
                "Could not extract pages from a password-protected PDF."
            )

        return [
            page.get_text()
            for page in doc
        ]

    except ValueError:
        raise

    except Exception as exc:
        raise ValueError(
            "Could not extract PDF pages. "
            "The file may be corrupted, unreadable, or unsupported."
        ) from exc

    finally:
        if doc is not None:
            doc.close()


def extract_docx_elements(
    file_bytes: bytes,
) -> list[dict]:
    """
    Extract addressable DOCX text elements for future DOCX redaction.

    This function is read-only.

    Each returned element identifies where the text came from:
        - header
        - paragraph
        - table_cell
        - footer

    Table-cell entries are paragraph-granular so future redaction can
    deterministically identify the exact paragraph within a cell.
    """
    try:
        document = Document(
            io.BytesIO(file_bytes)
        )

        elements: list[dict] = []

        # --------------------------------------------------
        # HEADERS
        # --------------------------------------------------
        for section_index, section in enumerate(
            document.sections
        ):
            for paragraph_index, paragraph in enumerate(
                section.header.paragraphs
            ):
                text = paragraph.text

                if text and text.strip():
                    elements.append(
                        {
                            "type": "header",
                            "section_index": section_index,
                            "paragraph_index": paragraph_index,
                            "text": text,
                        }
                    )

        # --------------------------------------------------
        # BODY PARAGRAPHS
        # --------------------------------------------------
        for paragraph_index, paragraph in enumerate(
            document.paragraphs
        ):
            text = paragraph.text

            if text and text.strip():
                elements.append(
                    {
                        "type": "paragraph",
                        "paragraph_index": paragraph_index,
                        "text": text,
                    }
                )

        # --------------------------------------------------
        # TABLE CELLS — PARAGRAPH GRANULAR
        # --------------------------------------------------
        for table_index, table in enumerate(
            document.tables
        ):
            for row_index, row in enumerate(
                table.rows
            ):
                for cell_index, cell in enumerate(
                    row.cells
                ):
                    for paragraph_index, paragraph in enumerate(
                        cell.paragraphs
                    ):
                        text = paragraph.text

                        if text and text.strip():
                            elements.append(
                                {
                                    "type": "table_cell",
                                    "table_index": table_index,
                                    "row_index": row_index,
                                    "cell_index": cell_index,
                                    "paragraph_index": paragraph_index,
                                    "text": text,
                                }
                            )

        # --------------------------------------------------
        # FOOTERS
        # --------------------------------------------------
        for section_index, section in enumerate(
            document.sections
        ):
            for paragraph_index, paragraph in enumerate(
                section.footer.paragraphs
            ):
                text = paragraph.text

                if text and text.strip():
                    elements.append(
                        {
                            "type": "footer",
                            "section_index": section_index,
                            "paragraph_index": paragraph_index,
                            "text": text,
                        }
                    )

        return elements

    except (
        PackageNotFoundError,
        KeyError,
        ValueError,
        OSError,
    ) as exc:
        raise ValueError(
            "Corrupted or unreadable DOCX file."
        ) from exc

    except Exception as exc:
        raise ValueError(
            "Corrupted or unreadable DOCX file."
        ) from exc